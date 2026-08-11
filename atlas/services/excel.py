"""Bounded, non-executing extraction for uploaded Excel workbooks."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

MAX_SHEETS = 8
MAX_ROWS_PER_SHEET = 200
MAX_COLUMNS = 40
MAX_CELLS = 5_000
MAX_CHARS = 80_000
MAX_CELL_CHARS = 500

EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}
EXCEL_MIMES = {
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.macroenabled.12",
}


def is_excel(filename: str | None, mime_type: str | None) -> bool:
    suffix = Path(filename or "").suffix.lower()
    return suffix in EXCEL_EXTENSIONS or (mime_type or "").lower() in EXCEL_MIMES


def extract_workbook(data: bytes, filename: str = "workbook.xlsx") -> str:
    """Return a compact TSV representation. Macros and formulas are never executed."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".xls":
        sheets = _read_xls(data)
    else:
        sheets = _read_openxml(data)
    return _render(filename, sheets)


def _read_openxml(data: bytes) -> list[tuple[str, int, int, Iterable[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    output = []
    try:
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            rows = (
                [cell.value for cell in row[:MAX_COLUMNS]]
                for row in sheet.iter_rows(max_row=MAX_ROWS_PER_SHEET, max_col=MAX_COLUMNS)
            )
            output.append((sheet.title, sheet.max_row or 0, sheet.max_column or 0, list(rows)))
    finally:
        workbook.close()
    return output


def _read_xls(data: bytes) -> list[tuple[str, int, int, Iterable[list[Any]]]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    output = []
    try:
        for sheet in workbook.sheets()[:MAX_SHEETS]:
            rows = [
                [sheet.cell_value(row, col) for col in range(min(sheet.ncols, MAX_COLUMNS))]
                for row in range(min(sheet.nrows, MAX_ROWS_PER_SHEET))
            ]
            output.append((sheet.name, sheet.nrows, sheet.ncols, rows))
    finally:
        workbook.release_resources()
    return output


def _render(filename: str, sheets: list[tuple[str, int, int, Iterable[list[Any]]]]) -> str:
    chunks = [f"Excel workbook: {Path(filename).name}", f"Sheets: {len(sheets)}"]
    cells = 0
    truncated = False

    for title, total_rows, total_columns, rows in sheets:
        chunks.append(f"\n--- Sheet: {title} ({total_rows} rows x {total_columns} columns) ---")
        emitted = 0
        for raw in rows:
            values = [_display(value) for value in raw]
            while values and values[-1] == "":
                values.pop()
            if not values:
                continue
            if cells + len(values) > MAX_CELLS:
                truncated = True
                break
            line = "\t".join(values)
            if sum(len(part) + 1 for part in chunks) + len(line) > MAX_CHARS:
                truncated = True
                break
            chunks.append(line)
            cells += len(values)
            emitted += 1
        if emitted < total_rows or total_columns > MAX_COLUMNS:
            truncated = True
        if truncated and (cells >= MAX_CELLS or sum(map(len, chunks)) >= MAX_CHARS - 1000):
            break

    if truncated:
        chunks.append("\n[Workbook preview truncated to safe processing limits.] ")
    chunks.append("[Formula cells use the workbook's last cached values; macros were not executed.]")
    return "\n".join(chunks)


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        text = f"{value:.10g}"
    else:
        text = str(value)
    return " ".join(text.replace("\t", " ").split())[:MAX_CELL_CHARS]
